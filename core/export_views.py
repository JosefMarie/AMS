import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.contrib.auth.decorators import login_required
from .models import Classroom, Attendance, CustomUser
from datetime import datetime, timedelta

@login_required
def export_attendance_excel(request):
    class_id = request.GET.get('class_id')
    start_date_str = request.GET.get('start_date')
    end_date_str = request.GET.get('end_date')

    if not class_id:
        return HttpResponse("Class ID is required.", status=400)

    classroom = get_object_or_404(Classroom, id=class_id)
    
    # Parse dates or use defaults
    try:
        start_date = datetime.strptime(start_date_str, '%Y-%m-%d').date() if start_date_str else None
        end_date = datetime.strptime(end_date_str, '%Y-%m-%d').date() if end_date_str else None
    except ValueError:
        return HttpResponse("Invalid date format. Use YYYY-MM-DD.", status=400)

    # Get students in the classroom
    students = CustomUser.objects.filter(student_profile__classroom=classroom, role=CustomUser.Role.STUDENT).select_related('student_profile').order_by('first_name', 'last_name')
    
    # Get attendance records
    attendances = Attendance.objects.filter(classroom=classroom)
    if start_date:
        attendances = attendances.filter(date__gte=start_date)
    if end_date:
        attendances = attendances.filter(date__lte=end_date)
        
    # Get unique dates
    dates = list(attendances.values_list('date', flat=True).distinct().order_by('date'))

    # Build attendance map: (student_id, date) -> status
    attendance_map = {}
    for att in attendances:
        attendance_map[(att.student_id, att.date)] = att.status

    # Create Excel workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Attendance - {classroom.name}"[:31] # Excel limits sheet name to 31 chars

    # Styles
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")

    # Title Row
    ws.merge_cells('A1:C1')
    ws['A1'] = f"Attendance Report: {classroom.name}"
    ws['A1'].font = Font(bold=True, size=14)
    
    if start_date and end_date:
        ws.merge_cells('A2:C2')
        ws['A2'] = f"Period: {start_date} to {end_date}"
    
    # Headers
    headers = ["Student ID", "First Name", "Last Name"]
    for d in dates:
        headers.append(d.strftime('%Y-%m-%d'))
    headers.extend(["Total Present", "Total Absent", "Total Late"])

    row_num = 4
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=row_num, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    # Set column widths
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 20
    for i in range(len(dates) + 3):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i + 4)].width = 12

    # Fill data
    row_num += 1
    for student in students:
        student_id = getattr(student.student_profile, 'student_id', student.username)
        ws.cell(row=row_num, column=1, value=student_id)
        ws.cell(row=row_num, column=2, value=student.first_name)
        ws.cell(row=row_num, column=3, value=student.last_name)
        
        present_count = 0
        absent_count = 0
        late_count = 0
        
        for i, d in enumerate(dates):
            col_num = 4 + i
            status = attendance_map.get((student.id, d), "-")
            
            if status == 'PRESENT':
                val = "P"
                present_count += 1
            elif status == 'ABSENT':
                val = "A"
                absent_count += 1
            elif status == 'LATE':
                val = "L"
                late_count += 1
            else:
                val = "-"
                
            cell = ws.cell(row=row_num, column=col_num, value=val)
            cell.alignment = center_align
            
            if val == "P":
                cell.font = Font(color="059669") # Green
            elif val == "A":
                cell.font = Font(color="DC2626") # Red
            elif val == "L":
                cell.font = Font(color="D97706") # Yellow

        ws.cell(row=row_num, column=4 + len(dates), value=present_count).alignment = center_align
        ws.cell(row=row_num, column=5 + len(dates), value=absent_count).alignment = center_align
        ws.cell(row=row_num, column=6 + len(dates), value=late_count).alignment = center_align
        
        row_num += 1

    # Prepare Response
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="Attendance_{classroom.name}.xlsx"'
    wb.save(response)
    return response
